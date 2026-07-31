"""Read dbg.xlsx and show structure"""
import openpyxl

wb = openpyxl.load_workbook(
    r"D:\GS\项目\FPGA\xcvu9p-flga2104-2L-e\fpga_pll_mash111_dtc_ssc\RTL\testbench\openclaw\dbg.xlsx",
    data_only=True  # show computed values, not formulas
)

for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n=== Sheet: {sn} ===")
    print(f"Dimensions: {ws.dimensions}")
    print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
    
    # Print all cells with content
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=False):
        row_data = []
        for cell in row:
            if cell.value is not None:
                row_data.append(f"{cell.coordinate}={cell.value}")
        if row_data:
            print(f"  {' | '.join(row_data)}")

wb.close()

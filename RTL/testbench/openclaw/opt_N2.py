"""Optimize I to minimize N25 in current dbg.xlsx"""
import openpyxl
import numpy as np

wb = openpyxl.load_workbook(
    r"D:\GS\项目\FPGA\xcvu9p-flga2104-2L-e\fpga_pll_mash111_dtc_ssc\RTL\testbench\openclaw\dbg.xlsx",
    data_only=True
)
ws = wb['Sheet1']

D = np.array([ws.cell(r, 4).value for r in range(2, 25)])  # rows 2-24
H = np.array([ws.cell(r, 8).value for r in range(2, 25)])

print(f"D rows: {len(D)}, H rows: {len(H)}")
print(f"D: {D[0]:.1f} ~ {D[-1]:.1f}")
print(f"H: {H[0]:.0f} ~ {H[-1]:.0f}")

# N[n] = J[n] - 2*J[n-1] + J[n-2]
#      = (D[n]-2*D[n-1]+D[n-2]) + I*(H[n]-2*H[n-1]+H[n-2])
# N starts from index 2 (row 4)
base = D[2:] - 2*D[1:-1] + D[:-2]
dH2 = H[2:] - 2*H[1:-1] + H[:-2]

print(f"\nbase (D 2nd diff): range {base.min():.1f} ~ {base.max():.1f}, mean={base.mean():.1f}")
print(f"dH2 (H 2nd diff): range {dH2.min():.0f} ~ {dH2.max():.0f}, mean={dH2.mean():.2f}")

# N25 = mean(base) + I * mean(dH2)
# For N25 = 0: I = -mean(base) / mean(dH2)
I_opt_N25 = -base.mean() / dH2.mean()
print(f"\nI for N25=0: I = {I_opt_N25:.6f}")
print(f"  = -({base.mean():.1f}) / ({dH2.mean():.2f})")

# Full scan
def N25_for_I(I):
    N = base + I * dH2
    return np.mean(N)

def std_N_for_I(I):
    N = base + I * dH2
    return np.std(N)

print(f"\n=== Scan ===")
best_std = (1e9, None)
for I_val in np.arange(-0.02, 0.02, 0.000001):
    n25 = N25_for_I(I_val)
    s = std_N_for_I(I_val)
    if s < best_std[0]:
        best_std = (s, I_val)

print(f"Min std(N): I = {best_std[1]:.6f}, std(N) = {best_std[0]:.4f}")

print(f"\n=== Key I values ===")
for I_val in [I_opt_N25, best_std[1], -0.005209, 0, 0.001, 0.002, 0.003, 1.0]:
    print(f"  I={I_val:+.6f}:  N25={N25_for_I(I_val):.4f}  std(N)={std_N_for_I(I_val):.4f}")

wb.close()

"""Analyze N column optimization (cycle-to-cycle jitter)"""
import openpyxl
import numpy as np

wb = openpyxl.load_workbook(
    r"D:\GS\项目\FPGA\xcvu9p-flga2104-2L-e\fpga_pll_mash111_dtc_ssc\RTL\testbench\openclaw\dbg.xlsx",
    data_only=True
)
ws = wb['Sheet1']

D = np.array([ws.cell(r, 4).value for r in range(3, 25)])
H = np.array([ws.cell(r, 8).value for r in range(3, 25)])

def compute_N_metrics(I):
    J = D + H * I
    M = np.diff(J)                     # periods
    N = np.diff(M)                     # cycle-to-cycle period diff
    return {
        'mean_N': np.mean(N),
        'std_N': np.std(N),
        'rms_N': np.sqrt(np.mean(N**2)),
        'N25': np.mean(N),             # = AVG(N)
        'max_N': np.max(np.abs(N)),
    }

# Scan I
print(f"Current I = -0.005209")
m = compute_N_metrics(-0.005209)
print(f"  mean_N = {m['mean_N']:.6f}")
print(f"  std_N  = {m['std_N']:.6f}")
print(f"  rms_N  = {m['rms_N']:.6f}")
print(f"  N25    = {m['N25']:.6f}")
print()

# Find I that minimizes |N25|, std_N, and rms_N
best_mean = (1e9, None)
best_std = (1e9, None)
best_rms = (1e9, None)

for I in np.arange(-0.05, 0.05, 0.000001):
    m = compute_N_metrics(I)
    if abs(m['N25']) < abs(best_mean[0]):
        best_mean = (m['N25'], I)
    if m['std_N'] < best_std[0]:
        best_std = (m['std_N'], I)
    if m['rms_N'] < best_rms[0]:
        best_rms = (m['rms_N'], I)

print("Optimal I values:")
print(f"  Min |N25| (mean diff):  I = {best_mean[1]:.6f}  N25 = {best_mean[0]:.10f}")
print(f"  Min std_N (variation):  I = {best_std[1]:.6f}  std_N = {best_std[0]:.6f}")
print(f"  Min rms_N (overall):    I = {best_rms[1]:.6f}  rms_N = {best_rms[0]:.6f}")
print()

# The user's value
print(f"Your value I = -0.005209:")
m = compute_N_metrics(-0.005209)
print(f"  N25 = {m['N25']:.10f}")
print(f"  std_N = {m['std_N']:.4f}")
print()

# Show a few candidate I values
print("Detailed comparison:")
for I_test in [-0.005209, best_mean[1], best_std[1], 0, -0.008, -0.005]:
    m = compute_N_metrics(I_test)
    print(f"  I={I_test:+.6f}:  N25={m['N25']:.6e}  std_N={m['std_N']:.4f}  rms_N={m['rms_N']:.4f}  max|N|={m['max_N']:.1f}")

wb.close()

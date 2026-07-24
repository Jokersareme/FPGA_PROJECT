"""Minimize max(|N|) = worst-case cycle-to-cycle jitter"""
import openpyxl
import numpy as np

wb = openpyxl.load_workbook(
    r"D:\GS\项目\FPGA\xcvu9p-flga2104-2L-e\fpga_pll_mash111_dtc_ssc\RTL\testbench\openclaw\dbg.xlsx",
    data_only=True
)
ws = wb['Sheet1']

D = np.array([ws.cell(r, 4).value for r in range(2, 25)])
H = np.array([ws.cell(r, 8).value for r in range(2, 25)])

base = D[2:] - 2*D[1:-1] + D[:-2]
dH2 = H[2:] - 2*H[1:-1] + H[:-2]

def max_abs_N(I):
    N = base + I * dH2
    return np.max(np.abs(N))

print(f"Searching for I that minimizes max(|N|)...")
print(f"base range: {base.min():.1f} ~ {base.max():.1f}")
print(f"dH2 range: {dH2.min():.0f} ~ {dH2.max():.0f}")
print()

best = (1e9, None)
for I_val in np.arange(-0.1, 0.1, 0.000001):
    m = max_abs_N(I_val)
    if m < best[0]:
        best = (m, I_val)

print(f"Min max(|N|) at I = {best[1]:.6f}")
print(f"  max(|N|) = {best[0]:.3f}")
print()

# Show N values at this optimal I
I_opt = best[1]
N_opt = base + I_opt * dH2
print(f"N values at I={I_opt:.6f}:")
for i, n in enumerate(N_opt):
    print(f"  N{i+4} = {n:.3f}")
print(f"  N25 = {np.mean(N_opt):.4f}")
print(f"  std(N) = {np.std(N_opt):.4f}")
print(f"  max|N| = {np.max(np.abs(N_opt)):.3f}")
print()

# Compare with key candidates
print("Comparison:")
for I_test in [0, 0.002, 0.004, 0.005, best[1], 0.005208, 0.006, 0.008, 0.009376, 0.01, 0.02]:
    m = max_abs_N(I_test)
    s = np.std(base + I_test * dH2)
    n25 = np.mean(base + I_test * dH2)
    print(f"  I={I_test:+.6f}:  max|N|={m:6.1f}  std(N)={s:6.2f}  N25={n25:6.4f}")

wb.close()

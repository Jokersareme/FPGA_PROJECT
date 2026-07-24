"""Find optimal J alignment (shift J relative to D)"""
import openpyxl
import numpy as np

wb = openpyxl.load_workbook(
    r"D:\GS\项目\FPGA\xcvu9p-flga2104-2L-e\fpga_pll_mash111_dtc_ssc\RTL\testbench\openclaw\dbg.xlsx",
    data_only=True
)
ws = wb['Sheet1']

D = np.array([ws.cell(r, 4).value for r in range(2, 23)])
J = np.array([ws.cell(r, 10).value for r in range(2, 23)])
L, N = 0.5, 0.0625

def compute(shift):
    if shift >= 0:
        Js, Ds = J[:len(J)-shift], D[shift:]
    else:
        s = -shift
        Js, Ds = J[s:], D[:len(D)-s]
    n = min(len(Js), len(Ds))
    if n < 4: return None
    Js, Ds = Js[:n], Ds[:n]
    M = (Js/4)*L + 256
    O = Ds + M*N
    R = np.diff(O)
    S = np.diff(R)
    return {'S23': np.mean(S), 'std_S': np.std(S), 'max_abs_S': np.max(np.abs(S)), 'n': len(S)}

print(f"{'Shift':>6} | {'S23':>10} | {'std_S':>8} | {'max|S|':>8} | {'n':>3}")
print("-" * 50)

results = []
for shift in range(-8, 13):
    m = compute(shift)
    if m:
        results.append((shift, m))
        print(f"  {shift:+4d} | {m['S23']:10.4f} | {m['std_S']:8.2f} | {m['max_abs_S']:8.1f} | {m['n']:3d}")

print()
if results:
    best_s23 = min(results, key=lambda x: abs(x[1]['S23']))
    best_max = min(results, key=lambda x: x[1]['max_abs_S'])
    best_std = min(results, key=lambda x: x[1]['std_S'])
    print(f"Best |S23|:  shift={best_s23[0]:+d}, S23={best_s23[1]['S23']:.4f}")
    print(f"Best max|S|: shift={best_max[0]:+d}, max|S|={best_max[1]['max_abs_S']:.2f}")
    print(f"Best std_S:  shift={best_std[0]:+d}, std_S={best_std[1]['std_S']:.4f}")

wb.close()

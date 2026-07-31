"""Find L (gain) that minimizes S23 (cycle-to-cycle jitter)"""
import openpyxl
import numpy as np

wb = openpyxl.load_workbook(
    r"D:\GS\项目\FPGA\xcvu9p-flga2104-2L-e\fpga_pll_mash111_dtc_ssc\RTL\testbench\openclaw\dbg.xlsx",
    data_only=True
)
ws = wb['Sheet1']

D = np.array([ws.cell(r, 4).value for r in range(2, 23)])  # D2:D22
J = np.array([ws.cell(r, 10).value for r in range(2, 23)]) # J2:J22 = mash eo
N = ws.cell(2, 14).value  # dtc2time gain

print(f"N (dtc2time gain) = {N}")

# M = dtc code = (J/4)*L + 256
# O = D + M * N = D + ((J/4)*L + 256) * N
# R = O[n] - O[n-1] = period after DTC
# S = R[n] - R[n-1] = cycle-to-cycle period variation

def compute_S23(L_val):
    M = (J / 4) * L_val + 256
    O = D + M * N
    R = np.diff(O)  # 20 values (R3:R22)
    S = np.diff(R)  # 19 values (S4:S22)
    return {
        'mean_S': np.mean(S),
        'std_S': np.std(S),
        'max_abs_S': np.max(np.abs(S)),
        'S23': np.mean(S),
    }

print(f"Current L=0.5: S23 = {compute_S23(0.5)['S23']:.4f}, max|S| = {compute_S23(0.5)['max_abs_S']:.1f}")
print()

# Find L that minimizes |S23|
best_mean = (1e9, None)
best_max = (1e9, None)
best_std = (1e9, None)

# Scan L from 0 to 5
for L_val in np.arange(0, 5.001, 0.0001):
    m = compute_S23(L_val)
    if abs(m['S23']) < abs(best_mean[0]):
        best_mean = (m['S23'], L_val)
    if m['max_abs_S'] < best_max[0]:
        best_max = (m['max_abs_S'], L_val)
    if m['std_S'] < best_std[0]:
        best_std = (m['std_S'], L_val)

print("Optimal L values:")
print(f"  Min |S23| (avg zero crossing): L = {best_mean[1]:.4f}  S23 = {best_mean[0]:.6f}")
print(f"  Min max|S| (worst-case):       L = {best_max[1]:.4f}  max|S| = {best_max[0]:.3f}")
print(f"  Min std_S (variation):          L = {best_std[1]:.4f}  std_S = {best_std[0]:.4f}")
print()

# Show comparison
print("Comparison:")
for L_test in [best_mean[1], best_max[1], best_std[1], 0, 0.5, 1, 1.5, 2, 2.5, 3, 4, 5]:
    m = compute_S23(L_test)
    if L_test >= 0:
        print(f"  L={L_test:6.4f}:  S23={m['S23']:8.4f}  std_S={m['std_S']:6.2f}  max|S|={m['max_abs_S']:6.1f}")

# Also show hardware gain equivalent
print(f"\nL -> hw_gain (L * 512):")
for L_test in [best_mean[1], best_max[1], best_std[1], 0.5, 1, 2]:
    print(f"  L={L_test:.4f} -> hw_gain={L_test*512:.1f}")

wb.close()
